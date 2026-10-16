"""Gate 0 — integrity check of the Happy Time curriculum workbook (read-only)."""
import hashlib, json, sys
import openpyxl
from openpyxl.utils import get_column_letter

path, manifest_path = sys.argv[1], sys.argv[2]
man = json.load(open(manifest_path, encoding="utf-8"))
results = []
info = []
import re

def check(name, expected, actual):
    ok = expected == actual
    results.append((name, expected, actual, ok))

raw = open(path, "rb").read()
check("sha256", man["authority"]["sha256"], hashlib.sha256(raw).hexdigest())
check("size_bytes", man["authority"]["size_bytes"], len(raw))

wb_f = openpyxl.load_workbook(path, data_only=False)  # formulas
wb_v = openpyxl.load_workbook(path, data_only=True)   # cached values
check("sheet_names", ["Malla ajustada", "Observaciones", "Progresión por dimensión"], wb_f.sheetnames)

def used_range(ws):
    max_r = max_c = 0
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and str(c.value).strip() != "":
                max_r = max(max_r, c.row); max_c = max(max_c, c.column)
    return f"A1:{get_column_letter(max_c)}{max_r}", max_r, max_c

for spec in man["authority_sheets"]:
    ws = wb_f[spec["sheet"]]
    rng, r, c = used_range(ws)
    check(f"{spec['sheet']} used range (content)", spec["range"], rng)
    info.append({"sheet": spec["sheet"], "cell_extent_incl_empty_styled": ws.dimensions, "content_range": rng})
    check(f"{spec['sheet']} rows", spec["rows"], r)
    check(f"{spec['sheet']} cols", spec["columns"], c)

ma = wb_f["Malla ajustada"]
nonempty_all = sum(1 for row in ma.iter_rows(min_row=1, max_row=21, max_col=33) for c in row if c.value is not None and str(c.value).strip() != "")
nonempty_exp = sum(1 for row in ma.iter_rows(min_row=2, max_row=21, max_col=33) for c in row if c.value is not None and str(c.value).strip() != "")
check("Malla nonempty all", 588, nonempty_all)
check("Malla nonempty experience rows", 555, nonempty_exp)

# headers + codes
headers = [ma.cell(1, j).value for j in range(1, 34)]
code_col = headers.index("Periodo / código") + 1
codes = []
for i in range(2, 22):
    m = re.search(r"\(([A-Z]\d?\.\d)\)\s*$", str(ma.cell(i, code_col).value or ""))
    codes.append(m.group(1) if m else None)
check("period codes (ordered)", man["certified_structure"]["period_codes"], codes)
check("experiences", 20, len(codes))
check("codes unique", 20, len(set(codes)))
level_labels = [ma.cell(i, 1).value for i in range(2, 22) if ma.cell(i, 1).value]
check("levels", 5, len(level_labels))
info.append({"level_labels_col_A": level_labels})
check("header count Malla", 33, sum(1 for h in headers if h))

obs = wb_f["Observaciones"]
obs_vals = [obs.cell(i, 1).value for i in range(1, 9)]
check("observations (rows 2-8 nonempty)", 7, sum(1 for v in obs_vals[1:] if v not in (None, "")))

pr = wb_f["Progresión por dimensión"]
prog = sum(1 for i in range(2, 7) for j in range(2, 8) if pr.cell(i, j).value not in (None, ""))
check("progression contents (B2:G6)", 30, prog)

# formulas / merged / hidden
formulas = [(ws.title, c.coordinate, c.value) for ws in wb_f for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
merged = {ws.title: [str(m) for m in ws.merged_cells.ranges] for ws in wb_f}
hidden = {ws.title: ws.sheet_state for ws in wb_f}

print(json.dumps({
    "checks": [{"check": n, "expected": e, "actual": a, "pass": ok} for n, e, a, ok in results],
    "all_pass": all(r[3] for r in results),
    "informative": info,
    "code_column": get_column_letter(code_col) if code_col else None,
    "headers": headers,
    "formulas": formulas[:50], "formula_count": len(formulas),
    "merged": merged, "sheet_state": hidden,
}, ensure_ascii=False, indent=1, default=str))
