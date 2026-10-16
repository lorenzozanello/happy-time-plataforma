"""
Validador independiente de fidelidad curricular (ADR-002, criterio F1-05).

Lee el Excel certificado con openpyxl (parser y lenguaje distintos a los del compilador TypeScript)
y verifica que la release compilada:
  1. reproduce cada celda de autoridad byte a byte;
  2. ATRIBUYE cada celda a la unidad correcta (id, tipo, campo, nivel, código, celda, rango);
  3. no omite ni duplica celdas;
  4. tiene estructuras coherentes (niveles, experiencias, progresión, observaciones, propósitos);
  5. tiene segmentos contiguos, con claves y etiquetas acordes a su literal.

Los mapas de encabezados y niveles están DUPLICADOS a propósito (no se leen de fields.ts ni de
levels.json): si alguien cambia esos archivos por error, esta verificación lo detecta.

Uso: python3 -I verify_independent.py curriculum/source/SOURCE.lock.json
Sale con código 1 ante cualquier discrepancia.
"""
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

# ── Mapas independientes ──────────────────────────────────────────────────────
LEVEL_FIELDS = {
    "Nivel": "nivel",
    "Ciclo normativo": "ciclo_normativo",
    "Equivalencia": "equivalencia",
    "Edad de entrada aprox.": "edad_entrada",
    "Edad de salida aprox.": "edad_salida",
    "Perfil de entrada sintetizado": "perfil_entrada",
    "Perfil de salida sintetizado": "perfil_salida",
}
PERIOD_HEADER = "Periodo / código"
TITLE_HEADER = "Experiencia integradora"
EXPERIENCE_FIELDS = {
    "Alcance de la experiencia integradora": "alcance",
    "Propósitos esenciales": "propositos",
    "Momento de Indagar": "momento_indagar",
    "Momento de Proyectar": "momento_proyectar",
    "Momento de Vivir la experiencia": "momento_vivir",
    "Momento de Valorar el proceso": "momento_valorar",
    "Juego: tema, competencia, avance y evidencia": "juego",
    "Exploración del medio: tema, competencia, avance y evidencia": "exploracion_medio",
    "Literatura: tema, competencia, avance y evidencia": "literatura",
    "Arte: tema, competencia, avance y evidencia": "arte",
    "Núcleos de experiencia": "nucleos",
    "Saberes y nociones movilizadas": "saberes",
    "Competencias de desarrollo": "competencias",
    "Evidencias pedagógicas integradas": "evidencias",
    "Diferenciación, DUA, apoyos razonables y PIAR cuando aplique": "dua_apoyos_piar",
    "Contexto Caribe": "contexto_caribe",
    "GIRE": "gire",
    "Procesos imprescindibles": "procesos_imprescindibles",
    "Ambientes saludables": "ambientes_saludables",
    "Indicadores de desarrollo": "indicadores",
    "Documentación pedagógica y alertas": "documentacion_alertas",
    "Eje de Educación Emocional": "educacion_emocional",
    "Alianza familia-escuela": "familia_escuela",
    "Observación y valoración": "observacion_valoracion",
}
DIMENSIONS = {
    "Cognitiva": "cognitiva",
    "Comunicativa": "comunicativa",
    "Socioemocional": "socioemocional",
    "Corporal": "corporal",
    "Estética": "estetica",
    "Ético-moral": "etico_moral",
}
LEVEL_BY_MALLA_LABEL = {
    "Step 1": "step-1",
    "Step 2": "step-2",
    "Step 3": "step-3",
    "Kínder": "kinder",
    "Transición": "transicion",
}
LEVEL_BY_PROG_LABEL = {
    "Step 1": "step-1",
    "Step 2": "step-2",
    "Step 3": "step-3",
    "Jardín/Kínder": "kinder",
    "Transición": "transicion",
}
LEVEL_BY_PREFIX = {"S1": "step-1", "S2": "step-2", "S3": "step-3", "K": "kinder", "T": "transicion"}
ROMANS = {"I": 1, "II": 2, "III": 3, "IV": 4}
PERIOD_RE = re.compile(r"^Periodo (I|II|III|IV) \(((?:S[123]|K|T)\.[1-4])\)$")

RECTORAS = {"juego", "exploracion_medio", "literatura", "arte"}
RECTORA_SEGMENTS = [("tema", "Tema"), ("competencia", "Competencia"), ("avance", "Avance"), ("evidencia", "Evidencia")]
SEGMENTED = RECTORAS | {"propositos", "documentacion_alertas"}
LOGRO_START = "En esta experiencia se logra cuando"
JS_WHITESPACE = re.compile(r"^[\s﻿]*$")


def utf16_len(s: str) -> int:
    return len(s.encode("utf-16-le")) // 2


def utf16_slice(s: str, start: int, end: int) -> str:
    return s.encode("utf-16-le")[start * 2 : end * 2].decode("utf-16-le")


class Checker:
    def __init__(self) -> None:
        self.errors: list[str] = []

    def eq(self, where: str, what: str, actual: object, expected: object) -> None:
        if actual != expected:
            self.errors.append(f"{where}: {what} = {actual!r}, se esperaba {expected!r}")


def merge_range_of(ws, coord: str) -> str | None:
    for m in ws.merged_cells.ranges:
        if coord in m:
            return m.coord
    return None


def expected_units(wb, lock, c: Checker) -> tuple[dict[str, dict], dict]:
    """Deriva desde el Excel la unidad esperada para cada celda de autoridad."""
    exp: dict[str, dict] = {}
    ma = wb[lock["malla"]["sheet"]]
    headers = [ma.cell(1, j).value for j in range(1, 34)]
    col_of = {h: i + 1 for i, h in enumerate(headers)}
    for h in [*LEVEL_FIELDS, PERIOD_HEADER, TITLE_HEADER, *EXPERIENCE_FIELDS]:
        if headers.count(h) != 1:
            c.errors.append(f"Encabezado '{h}' ausente o duplicado en Malla")
    if c.errors:
        return exp, {}

    structure = {"levels": [], "experiences": []}
    block = lock["malla"]["level_block_rows"]
    for b in range(5):
        top = 2 + b * block
        label = ma.cell(top, col_of["Nivel"]).value
        lid = LEVEL_BY_MALLA_LABEL.get(label)
        if lid is None:
            c.errors.append(f"Nivel desconocido '{label}' en A{top}")
            continue
        level = {"id": lid, "label": label, "codes": [], "fields": {}}
        for header, key in LEVEL_FIELDS.items():
            col = col_of[header]
            coord = f"{get_column_letter(col)}{top}"
            uid = f"L:{lid}#{key}"
            level["fields"][key] = uid
            exp[uid] = {
                "kind": "level_field", "fieldKey": key, "levelId": lid, "code": None,
                "sheet": ma.title, "cell": coord, "header": header,
                "range": merge_range_of(ma, coord), "text": ma[coord].value,
            }
            expected_range = f"{coord}:{get_column_letter(col)}{top + block - 1}"
            c.eq(uid, "rango combinado en el Excel", exp[uid]["range"], expected_range)
        for row in range(top, top + block):
            period = ma.cell(row, col_of[PERIOD_HEADER]).value
            m = PERIOD_RE.match(period or "")
            if not m:
                c.errors.append(f"Periodo ilegible en fila {row}")
                continue
            roman, code = m.group(1), m.group(2)
            c.eq(code, "nivel según prefijo", LEVEL_BY_PREFIX[code.split(".")[0]], lid)
            c.eq(code, "periodo romano vs código", ROMANS[roman], int(code.split(".")[1]))
            level["codes"].append(code)
            exper = {
                "code": code, "levelId": lid, "periodIndex": ROMANS[roman], "periodLiteral": period,
                "periodCell": f"{get_column_letter(col_of[PERIOD_HEADER])}{row}", "fields": {},
            }
            tcoord = f"{get_column_letter(col_of[TITLE_HEADER])}{row}"
            exp[f"{code}#titulo"] = {
                "kind": "experience_title", "fieldKey": "titulo", "levelId": lid, "code": code,
                "sheet": ma.title, "cell": tcoord, "header": TITLE_HEADER, "range": None,
                "text": ma[tcoord].value,
            }
            for header, key in EXPERIENCE_FIELDS.items():
                coord = f"{get_column_letter(col_of[header])}{row}"
                uid = f"{code}#{key}"
                exper["fields"][key] = uid
                exp[uid] = {
                    "kind": "experience_field", "fieldKey": key, "levelId": lid, "code": code,
                    "sheet": ma.title, "cell": coord, "header": header, "range": None,
                    "text": ma[coord].value,
                }
            structure["experiences"].append(exper)
        structure["levels"].append(level)

    ob = wb[lock["observaciones"]["sheet"]]
    for i in range(1, lock["observaciones"]["count"] + 1):
        coord = f"A{i + 1}"
        exp[f"OBS-{i}"] = {
            "kind": "observation", "fieldKey": "observacion", "levelId": None, "code": None,
            "sheet": ob.title, "cell": coord, "header": ob["A1"].value, "range": None,
            "text": ob[coord].value,
        }

    pr = wb[lock["progresion"]["sheet"]]
    pr_headers = [pr.cell(1, j).value for j in range(1, 8)]
    structure["progression"] = []
    structure["progLabels"] = {}
    for row in range(2, 7):
        label = pr.cell(row, 1).value
        lid = LEVEL_BY_PROG_LABEL.get(label)
        if lid is None:
            c.errors.append(f"Nivel de progresión desconocido '{label}' en A{row}")
            continue
        structure["progLabels"][lid] = label
        for j in range(2, 8):
            header = pr_headers[j - 1]
            dim = DIMENSIONS.get(header)
            if dim is None:
                c.errors.append(f"Dimensión desconocida '{header}'")
                continue
            coord = f"{get_column_letter(j)}{row}"
            uid = f"PROG:{lid}#{dim}"
            structure["progression"].append((lid, dim, uid))
            exp[uid] = {
                "kind": "progression", "fieldKey": dim, "levelId": lid, "code": None,
                "sheet": pr.title, "cell": coord, "header": header, "range": None,
                "text": pr[coord].value,
            }
    return exp, structure


def check_units(data, exp: dict[str, dict], release: str, c: Checker) -> None:
    units = {u["unitId"]: u for u in data["units"]}
    c.eq("release", "unidades totales", len(data["units"]), len(exp))
    c.eq("release", "unitId únicos", len(units), len(data["units"]))
    for uid in exp.keys() - units.keys():
        c.errors.append(f"{uid}: falta en la release (celda {exp[uid]['sheet']}!{exp[uid]['cell']} sin cubrir)")
    for uid in units.keys() - exp.keys():
        c.errors.append(f"{uid}: unidad sin celda de autoridad correspondiente")
    for uid in exp.keys() & units.keys():
        e, u = exp[uid], units[uid]
        if not isinstance(e["text"], str):
            c.errors.append(f"{uid}: la celda esperada no contiene texto")
        c.eq(uid, "text", u["text"], e["text"])
        c.eq(uid, "kind", u["kind"], e["kind"])
        c.eq(uid, "fieldKey", u["fieldKey"], e["fieldKey"])
        c.eq(uid, "levelId", u.get("levelId"), e["levelId"])
        c.eq(uid, "code", u.get("code"), e["code"])
        c.eq(uid, "source.sheet", u["source"]["sheet"], e["sheet"])
        c.eq(uid, "source.cell", u["source"]["cell"], e["cell"])
        c.eq(uid, "source.header", u["source"]["header"], e["header"])
        c.eq(uid, "source.range", u["source"].get("range"), e["range"])
        c.eq(uid, "headerLiteral", u["headerLiteral"], e["header"])
        c.eq(uid, "authority", u["authority"], "official_curriculum")
        c.eq(uid, "release", u["release"], release)
        check_segments(u, c)


def check_segments(u: dict, c: Checker) -> None:
    uid, key, text, segs = u["unitId"], u["fieldKey"], u["text"], u.get("segments")
    if u["kind"] != "experience_field" or key not in SEGMENTED:
        if segs is not None:
            c.errors.append(f"{uid}: tiene segmentos y no debería")
        return
    if not segs:
        c.errors.append(f"{uid}: le faltan segmentos")
        return
    pos = 0
    for s in segs:
        if s["start"] != pos:
            c.errors.append(f"{uid}: segmento {s['key']} no contiguo")
        if not (s["start"] <= s["valueStart"] < s["valueEnd"] <= s["end"]):
            c.errors.append(f"{uid}: rangos inválidos o valor vacío en {s['key']}")
        suffix = utf16_slice(text, s["valueEnd"], s["end"])
        if not JS_WHITESPACE.match(suffix):
            c.errors.append(f"{uid}: sufijo con contenido en {s['key']}: {suffix!r}")
        pos = s["end"]
    if pos != utf16_len(text):
        c.errors.append(f"{uid}: los segmentos no cubren el texto completo")

    prefixes = [utf16_slice(text, s["start"], s["valueStart"]) for s in segs]
    keys_labels = [(s["key"], s["label"]) for s in segs]
    if key in RECTORAS:
        c.eq(uid, "claves de segmento", keys_labels, RECTORA_SEGMENTS)
        c.eq(uid, "prefijos", prefixes, [f"{label}: " for _, label in RECTORA_SEGMENTS])
    elif key == "documentacion_alertas":
        c.eq(uid, "claves de segmento", keys_labels, [("documentacion", "Documentación"), ("alertas", "Alertas")])
        c.eq(uid, "prefijos", prefixes, ["", "Alertas: "])
    else:
        expected = []
        for n in (1, 2, 3):
            expected += [(f"proposito_{n}", f"Propósito {n}"), (f"logro_{n}", "En esta experiencia")]
        c.eq(uid, "claves de segmento", keys_labels, expected)
        c.eq(uid, "prefijos", prefixes, ["1. ", "", "2. ", "", "3. ", ""])
        for s in segs:
            if s["key"].startswith("logro_"):
                value = utf16_slice(text, s["valueStart"], s["valueEnd"])
                if not value.startswith(LOGRO_START):
                    c.errors.append(f"{uid}: {s['key']} no empieza por el marcador de logro")


def check_structures(data, structure, wb, lock, c: Checker) -> None:
    c.eq("levels", "ids en orden", [l["id"] for l in data["levels"]], [l["id"] for l in structure["levels"]])
    for got, want in zip(data["levels"], structure["levels"]):
        lid = want["id"]
        c.eq(lid, "labelLiteral", got["labelLiteral"], want["label"])
        c.eq(lid, "progresionLabelLiteral", got["progresionLabelLiteral"], structure["progLabels"].get(lid))
        c.eq(lid, "periodCodes", got["periodCodes"], want["codes"])
        c.eq(lid, "fieldUnitIds", got["fieldUnitIds"], want["fields"])
        c.eq(lid, "codePrefix", got["codePrefix"], want["codes"][0].split(".")[0] if want["codes"] else None)

    c.eq("experiences", "códigos en orden", [e["code"] for e in data["experiences"]], lock["malla"]["period_codes"])
    for got, want in zip(data["experiences"], structure["experiences"]):
        code = want["code"]
        c.eq(code, "code", got["code"], code)
        c.eq(code, "levelId", got["levelId"], want["levelId"])
        c.eq(code, "periodIndex", got["periodIndex"], want["periodIndex"])
        c.eq(code, "periodLiteral", got["periodLiteral"], want["periodLiteral"])
        c.eq(code, "periodSource.cell", got["periodSource"]["cell"], want["periodCell"])
        c.eq(code, "periodSource.header", got["periodSource"]["header"], PERIOD_HEADER)
        c.eq(code, "titleUnitId", got["titleUnitId"], f"{code}#titulo")
        c.eq(code, "fieldUnitIds", got["fieldUnitIds"], want["fields"])

    c.eq("observationUnitIds", "lista", data["observationUnitIds"], [f"OBS-{i}" for i in range(1, 8)])
    c.eq("dimensions", "lista", [(d["key"], d["headerLiteral"]) for d in data["dimensions"]],
         [(v, k) for k, v in DIMENSIONS.items()])
    got_prog = sorted((p["levelId"], p["dimensionKey"], p["unitId"]) for p in data["progression"])
    c.eq("progression", "mapa nivel × dimensión", got_prog, sorted(structure["progression"]))

    units = {u["unitId"]: u for u in data["units"]}
    proposito_ids = [f"{e['code']}#propositos" for e in structure["experiences"]]
    for p in data["purposes"]:
        where = f"purposes[{p['n']}]"
        c.eq(where, "segmentKey", p["segmentKey"], f"proposito_{p['n']}")
        c.eq(where, "unitIds", p["unitIds"], proposito_ids)
        for uid in p["unitIds"]:
            u = units.get(uid)
            seg = next((s for s in (u or {}).get("segments") or [] if s["key"] == p["segmentKey"]), None)
            if not u or not seg:
                c.errors.append(f"{where}: {uid} sin segmento {p['segmentKey']}")
                continue
            c.eq(where, f"texto en {uid}", utf16_slice(u["text"], seg["valueStart"], seg["valueEnd"]), p["text"])

    # Toda celda no vacía de autoridad queda cubierta (unidades + literales de periodo + etiquetas).
    covered = {(u["source"]["sheet"], u["source"]["cell"]) for u in data["units"]}
    covered |= {(e["periodSource"]["sheet"], e["periodSource"]["cell"]) for e in data["experiences"]}
    ma, ob, pr = lock["malla"]["sheet"], lock["observaciones"]["sheet"], lock["progresion"]["sheet"]
    covered |= {(ma, f"{get_column_letter(j)}1") for j in range(1, 34)}
    covered |= {(pr, f"{get_column_letter(j)}1") for j in range(1, 8)}
    covered |= {(pr, f"A{r}") for r in range(2, 7)} | {(ob, "A1")}
    for sheet in lock["sheets"]:
        for row in wb[sheet].iter_rows():
            for cell in row:
                if cell.value not in (None, "") and (sheet, cell.coordinate) not in covered:
                    c.errors.append(f"{sheet}!{cell.coordinate} con contenido no cubierto")


def main() -> int:
    lock_path = Path(sys.argv[1]).resolve()
    root = lock_path.parents[2]
    lock = json.loads(lock_path.read_text(encoding="utf-8"))
    release = lock["release"]
    data = json.loads((root / "curriculum" / "build" / f"curriculum.{release}.json").read_text(encoding="utf-8"))
    c = Checker()

    raw = (root / lock["file"]).read_bytes()
    c.eq("fuente", "SHA-256", hashlib.sha256(raw).hexdigest(), lock["sha256"])
    c.eq("release", "source.sha256", data["source"]["sha256"], lock["sha256"])
    c.eq("release", "release", data["release"], release)

    wb = openpyxl.load_workbook(root / lock["file"], data_only=False)
    c.eq("fuente", "hojas", wb.sheetnames, lock["sheets"])
    if not c.errors:
        exp, structure = expected_units(wb, lock, c)
        if not c.errors:
            check_units(data, exp, release, c)
            check_structures(data, structure, wb, lock, c)

    counts = data["counts"]
    expected_counts = {"units": 572, "experienceFields": 480, "experienceTitles": 20, "levelFields": 35, "observations": 7, "progression": 30}
    for k, v in expected_counts.items():
        c.eq("counts", k, counts.get(k), v)

    report = {
        "release": release,
        "units_checked": len(data["units"]),
        "cells_covered": len(data["units"]) + len(data["experiences"]),
        "errors": c.errors,
        "pass": not c.errors,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not c.errors else 1


if __name__ == "__main__":
    sys.exit(main())
